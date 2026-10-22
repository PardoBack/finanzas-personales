import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_multiuser_excel(filename="PardoFinanzas_Multiusuario_Maestro.xlsx"):
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    navy_dark = "090D14"
    card_header = "121722"
    emerald_accent = "10B981"
    emerald_light = "ECFDF5"
    cyan_accent = "0284C7"
    cyan_light = "F0F9FF"
    amber_accent = "D97706"
    rose_accent = "E11D48"
    border_gray = "CBD5E1"
    zebra_even = "F8FAFC"
    white = "FFFFFF"

    font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="94A3B8")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_kpi_num = Font(name="Calibri", size=18, bold=True, color=navy_dark)
    font_kpi_label = Font(name="Calibri", size=10, bold=True, color="64748B")
    font_body = Font(name="Calibri", size=11, color="0F172A")
    font_body_bold = Font(name="Calibri", size=11, bold=True, color="0F172A")
    font_status_paid = Font(name="Calibri", size=10, bold=True, color="065F46")
    font_status_pend = Font(name="Calibri", size=10, bold=True, color="92400E")

    fill_banner = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
    fill_header = PatternFill(start_color=card_header, end_color=card_header, fill_type="solid")
    fill_header_emerald = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    fill_zebra = PatternFill(start_color=zebra_even, end_color=zebra_even, fill_type="solid")
    fill_kpi_card = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_paid = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    fill_pend = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=border_gray),
        right=Side(style='thin', color=border_gray),
        top=Side(style='thin', color=border_gray),
        bottom=Side(style='thin', color=border_gray)
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    num_format_currency = '"$"#,##0'
    num_format_pct = '0.0%'
    num_format_date = 'yyyy-mm-dd'

    # =========================================================
    # HOJA 1: USUARIOS
    # =========================================================
    ws1 = wb.create_sheet(title="01_USUARIOS")
    ws1.views.sheetView[0].showGridLines = True

    # Banner
    ws1.merge_cells("A1:G2")
    cell_b1 = ws1["A1"]
    cell_b1.value = "  PARDOFINANZAS • DIRECTORIO MULTIUSUARIO DE PARTICIPANTES"
    cell_b1.font = font_title
    cell_b1.fill = fill_banner
    cell_b1.alignment = align_left

    headers_u = ["ID_USUARIO", "NOMBRE_COMPLETO", "ALIAS", "ROL", "SALDO_INICIAL_COP", "EMAIL_NOTIFICACIONES", "ESTADO"]
    for col_idx, h in enumerate(headers_u, 1):
        c = ws1.cell(row=4, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = thin_border

    usuarios_data = [
        ("U-001", "Marlon Muñoz Pardo", "Marlon", "Administrador (Pro)", 1065366, "marlon@pardofinanzas.com", "ACTIVO"),
        ("U-002", "Carlos Restrepo", "Carlos", "Miembro", 500000, "carlos.r@gmail.com", "ACTIVO"),
        ("U-003", "Laura Morales", "Laura", "Miembro", 750000, "laura.m@hotmail.com", "ACTIVO"),
        ("U-004", "Andrés Gómez", "Andrés", "Miembro", 620000, "andres.g@gmail.com", "ACTIVO")
    ]

    for row_idx, u in enumerate(usuarios_data, 5):
        for col_idx, val in enumerate(u, 1):
            c = ws1.cell(row=row_idx, column=col_idx, value=val)
            c.font = font_body
            c.border = thin_border
            if col_idx == 1:
                c.alignment = align_center
                c.font = font_body_bold
            elif col_idx in (2, 3, 6):
                c.alignment = align_left
            elif col_idx == 4:
                c.alignment = align_center
            elif col_idx == 5:
                c.alignment = align_right
                c.number_format = num_format_currency
            elif col_idx == 7:
                c.alignment = align_center
                c.fill = fill_paid
                c.font = font_status_paid

    # =========================================================
    # HOJA 2: MOVIMIENTOS
    # =========================================================
    ws2 = wb.create_sheet(title="02_MOVIMIENTOS")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:K2")
    c_m1 = ws2["A1"]
    c_m1.value = "  LIBRO MAESTRO DE MOVIMIENTOS FINANCIEROS MULTIUSUARIO"
    c_m1.font = font_title
    c_m1.fill = fill_banner
    c_m1.alignment = align_left

    headers_m = [
        "ID_MOV", "ID_USUARIO", "USUARIO", "FECHA", "TIPO",
        "CATEGORIA", "CONCEPTO_DESCRIPCION", "VALOR_COP",
        "MEDIO_PAGO", "COMERCIO_DESTINO", "ESTADO"
    ]

    for col_idx, h in enumerate(headers_m, 1):
        c = ws2.cell(row=4, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header_emerald
        c.alignment = align_center
        c.border = thin_border

    sample_movs = [
        ("M-1001", "U-001", "2026-10-01", "Ingreso", "Salario", "Quincena Nómina Colanta", 1850000, "AyC Colanta", "Cooperativa Colanta", "CONFIRMADO"),
        ("M-1002", "U-001", "2026-10-02", "Gasto", "Alimentación", "Mercado Éxito mensual", 320000, "Nequi", "Éxito Colombia", "CONFIRMADO"),
        ("M-1003", "U-001", "2026-10-03", "Gasto", "Transporte/Moto", "Tanqueada Gasolina Extra", 35000, "Efectivo", "Estación Terpel", "CONFIRMADO"),
        ("M-1004", "U-001", "2026-10-15", "Gasto", "Servicios", "Factura EPM Energía", 197000, "AyC Colanta", "EPM Medellín", "CONFIRMADO"),
        ("M-1005", "U-002", "2026-10-01", "Ingreso", "Salario", "Salario Mensual Carlos", 2200000, "Bancolombia", "Empresa Tech", "CONFIRMADO"),
        ("M-1006", "U-002", "2026-10-04", "Gasto", "Alimentación", "Almuerzo Restaurante", 28000, "Nequi", "Restaurante Central", "CONFIRMADO"),
        ("M-1007", "U-003", "2026-10-01", "Ingreso", "Salario", "Honorarios Laura", 1600000, "Bancolombia", "Consultoría SAS", "CONFIRMADO"),
        ("M-1008", "U-003", "2026-10-06", "Gasto", "Salud/Gym", "Mensualidad Smart Fit", 99000, "Tarjeta Débito", "Smart Fit", "CONFIRMADO"),
        ("M-1009", "U-004", "2026-10-01", "Ingreso", "Salario", "Pago Quincenal Andrés", 1400000, "Nequi", "Comercio Global", "CONFIRMADO"),
        ("M-1010", "U-004", "2026-10-08", "Gasto", "Ocio", "Salida Cine y Cena", 65000, "Efectivo", "Cine Colombia", "CONFIRMADO")
    ]

    for r_idx, row_data in enumerate(sample_movs, 5):
        id_mov, id_usr, fecha, tipo, cat, desc, val, medio, com, estado = row_data

        ws2.cell(row=r_idx, column=1, value=id_mov).alignment = align_center
        ws2.cell(row=r_idx, column=2, value=id_usr).alignment = align_center
        # Formula VLOOKUP to sheet 01_USUARIOS compatible with Google Sheets
        ws2.cell(row=r_idx, column=3, value=f"=IFERROR(VLOOKUP(B{r_idx}, '01_USUARIOS'!$A$5:$B$20, 2, FALSE), \"-\")").alignment = align_left
        ws2.cell(row=r_idx, column=4, value=fecha).alignment = align_center
        ws2.cell(row=r_idx, column=5, value=tipo).alignment = align_center
        ws2.cell(row=r_idx, column=6, value=cat).alignment = align_left
        ws2.cell(row=r_idx, column=7, value=desc).alignment = align_left
        
        c_val = ws2.cell(row=r_idx, column=8, value=val)
        c_val.alignment = align_right
        c_val.number_format = num_format_currency

        ws2.cell(row=r_idx, column=9, value=medio).alignment = align_center
        ws2.cell(row=r_idx, column=10, value=com).alignment = align_left
        
        c_est = ws2.cell(row=r_idx, column=11, value=estado)
        c_est.alignment = align_center
        c_est.fill = fill_paid if estado == "CONFIRMADO" else fill_pend
        c_est.font = font_status_paid if estado == "CONFIRMADO" else font_status_pend

        for col_i in range(1, 12):
            ws2.cell(row=r_idx, column=col_i).border = thin_border
            if col_i != 11:
                ws2.cell(row=r_idx, column=col_i).font = font_body

    # =========================================================
    # HOJA 3: CONTROL DE FIJOS
    # =========================================================
    ws3 = wb.create_sheet(title="03_CONTROL_FIJOS")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:I2")
    c_f1 = ws3["A1"]
    c_f1.value = "  CONTROL DE GASTOS FIJOS (QUINCENA 1 - DÍA 15 vs FIN DE MES - DÍA 30)"
    c_f1.font = font_title
    c_f1.fill = fill_banner
    c_f1.alignment = align_left

    headers_f = [
        "ID_FIJO", "ID_USUARIO", "USUARIO", "CONCEPTO_FIJO",
        "CATEGORIA", "CUOTA_MENSUAL", "DIA_PAGO", "QUINCENA", "ESTADO_MES"
    ]

    for col_idx, h in enumerate(headers_f, 1):
        c = ws3.cell(row=4, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = thin_border

    sample_fijos = [
        ("F-01", "U-001", "EPM Energía", "Servicios", 197000, 15, "Pagado"),
        ("F-02", "U-001", "Plan Celular Tigo", "Servicios", 55000, 15, "Pagado"),
        ("F-03", "U-001", "Seguro Moto SURA", "Transporte", 135000, 15, "Pendiente"),
        ("F-04", "U-001", "Arriendo Medellín", "Vivienda", 920000, 30, "Pagado"),
        ("F-05", "U-001", "Internet Fibra Óptica", "Servicios", 105000, 30, "Pendiente"),
        ("F-06", "U-002", "Arriendo Habitación", "Vivienda", 650000, 30, "Pendiente"),
        ("F-07", "U-002", "Gimnasio", "Salud", 85000, 15, "Pagado"),
        ("F-08", "U-003", "Crédito Estudio", "Educación", 320000, 15, "Pagado"),
        ("F-09", "U-004", "Plan Móvil Claro", "Servicios", 45000, 30, "Pendiente")
    ]

    for r_idx, (fid, uid, conc, cat, cuota, dia, est) in enumerate(sample_fijos, 5):
        ws3.cell(row=r_idx, column=1, value=fid).alignment = align_center
        ws3.cell(row=r_idx, column=2, value=uid).alignment = align_center
        ws3.cell(row=r_idx, column=3, value=f"=IFERROR(VLOOKUP(B{r_idx}, '01_USUARIOS'!$A$5:$B$20, 2, FALSE), \"-\")").alignment = align_left
        ws3.cell(row=r_idx, column=4, value=conc).alignment = align_left
        ws3.cell(row=r_idx, column=5, value=cat).alignment = align_left

        c_q = ws3.cell(row=r_idx, column=6, value=cuota)
        c_q.alignment = align_right
        c_q.number_format = num_format_currency

        ws3.cell(row=r_idx, column=7, value=dia).alignment = align_center
        ws3.cell(row=r_idx, column=8, value=f'=IF(G{r_idx}<=15, "Quincena 1 (Día 15)", "Fin de Mes (Día 30)")').alignment = align_center

        c_e = ws3.cell(row=r_idx, column=9, value=est)
        c_e.alignment = align_center
        c_e.fill = fill_paid if est == "Pagado" else fill_pend
        c_e.font = font_status_paid if est == "Pagado" else font_status_pend

        for col_i in range(1, 10):
            ws3.cell(row=r_idx, column=col_i).border = thin_border
            if col_i != 9:
                ws3.cell(row=r_idx, column=col_i).font = font_body

    # =========================================================
    # HOJA 4: DASHBOARD MULTIUSUARIO
    # =========================================================
    ws4 = wb.create_sheet(title="04_DASHBOARD_RESUMEN")
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells("A1:H2")
    c_d1 = ws4["A1"]
    c_d1.value = "  PANEL CONSOLIDADO MULTIUSUARIO • FORMULAS 100% GOOGLE SHEETS"
    c_d1.font = font_title
    c_d1.fill = fill_banner
    c_d1.alignment = align_left

    # KPIs Globales (Cards)
    kpis = [
        ("A4:B4", "A5:B5", "INGRESOS TOTALES GRUPO", '=SUMIF(\'02_MOVIMIENTOS\'!E:E, "Ingreso", \'02_MOVIMIENTOS\'!H:H)', emerald_accent),
        ("C4:D4", "C5:D5", "GASTOS TOTALES GRUPO", '=SUMIF(\'02_MOVIMIENTOS\'!E:E, "Gasto", \'02_MOVIMIENTOS\'!H:H)', rose_accent),
        ("E4:F4", "E5:F5", "SUPERÁVIT NETO COLECTIVO", '=A5-C5', cyan_accent),
        ("G4:H4", "G5:H5", "TASA DE AHORRO COLECTIVA", '=IFERROR(E5/A5, 0)', amber_accent)
    ]

    for label_range, val_range, title, formula, color_accent in kpis:
        ws4.merge_cells(label_range)
        l_cell = ws4[label_range.split(":")[0]]
        l_cell.value = title
        l_cell.font = font_kpi_label
        l_cell.alignment = align_center
        l_cell.fill = fill_kpi_card

        ws4.merge_cells(val_range)
        v_cell = ws4[val_range.split(":")[0]]
        v_cell.value = formula
        v_cell.font = Font(name="Calibri", size=16, bold=True, color=color_accent)
        v_cell.alignment = align_center
        v_cell.fill = fill_kpi_card
        if "TASA" in title:
            v_cell.number_format = num_format_pct
        else:
            v_cell.number_format = num_format_currency

    # Tabla Desglose por Usuario
    ws4.cell(row=8, column=1, value="DESGLOSE INDIVIDUAL POR PARTICIPANTE").font = Font(name="Calibri", size=13, bold=True, color=navy_dark)

    headers_ind = [
        "ID_USUARIO", "NOMBRE_PARTICIPANTE", "SALDO_INICIAL",
        "INGRESOS_MES", "GASTOS_MES", "BALANCE_NETO",
        "TASA_AHORRO", "SALDO_ACTUAL_DISPONIBLE"
    ]

    for c_i, h in enumerate(headers_ind, 1):
        c = ws4.cell(row=10, column=c_i, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = thin_border

    for i in range(1, 5):
        r = 10 + i
        uid = f"U-00{i}"
        ws4.cell(row=r, column=1, value=uid).alignment = align_center
        ws4.cell(row=r, column=2, value=f"=IFERROR(VLOOKUP(A{r}, '01_USUARIOS'!$A$5:$B$20, 2, FALSE), \"-\")").alignment = align_left
        
        # Saldo Inicial
        c_s = ws4.cell(row=r, column=3, value=f"=IFERROR(VLOOKUP(A{r}, '01_USUARIOS'!$A$5:$E$20, 5, FALSE), 0)")
        c_s.alignment = align_right
        c_s.number_format = num_format_currency

        # Ingresos Mes
        c_ing = ws4.cell(row=r, column=4, value=f"=SUMIFS('02_MOVIMIENTOS'!$H:$H, '02_MOVIMIENTOS'!$B:$B, A{r}, '02_MOVIMIENTOS'!$E:$E, \"Ingreso\")")
        c_ing.alignment = align_right
        c_ing.number_format = num_format_currency

        # Gastos Mes
        c_gas = ws4.cell(row=r, column=5, value=f"=SUMIFS('02_MOVIMIENTOS'!$H:$H, '02_MOVIMIENTOS'!$B:$B, A{r}, '02_MOVIMIENTOS'!$E:$E, \"Gasto\")")
        c_gas.alignment = align_right
        c_gas.number_format = num_format_currency

        # Balance Neto
        c_bal = ws4.cell(row=r, column=6, value=f"=D{r}-E{r}")
        c_bal.alignment = align_right
        c_bal.number_format = num_format_currency

        # Tasa Ahorro
        c_tasa = ws4.cell(row=r, column=7, value=f"=IFERROR(F{r}/D{r}, 0)")
        c_tasa.alignment = align_center
        c_tasa.number_format = num_format_pct

        # Saldo Actual Disponible = Saldo Inicial + Balance Neto
        c_disp = ws4.cell(row=r, column=8, value=f"=C{r}+F{r}")
        c_disp.alignment = align_right
        c_disp.number_format = num_format_currency
        c_disp.font = font_body_bold

        for col_k in range(1, 9):
            ws4.cell(row=r, column=col_k).border = thin_border
            if col_k != 8:
                ws4.cell(row=r, column=col_k).font = font_body

    # Fila de Totales
    r_tot = 15
    ws4.cell(row=r_tot, column=1, value="TOTALES").alignment = align_center
    ws4.cell(row=r_tot, column=1).font = font_body_bold
    ws4.cell(row=r_tot, column=2, value="Todos los participantes").font = font_body_bold

    for col_k in [3, 4, 5, 6, 8]:
        col_letter = get_column_letter(col_k)
        c_t = ws4.cell(row=r_tot, column=col_k, value=f"=SUM({col_letter}11:{col_letter}14)")
        c_t.font = font_body_bold
        c_t.alignment = align_right
        c_t.number_format = num_format_currency

    c_tasa_tot = ws4.cell(row=r_tot, column=7, value=f"=IFERROR(F{r_tot}/D{r_tot}, 0)")
    c_tasa_tot.font = font_body_bold
    c_tasa_tot.alignment = align_center
    c_tasa_tot.number_format = num_format_pct

    for col_k in range(1, 9):
        ws4.cell(row=r_tot, column=col_k).border = thin_border
        ws4.cell(row=r_tot, column=col_k).fill = fill_zebra

    # =========================================================
    # HOJA 5: PARAMETROS
    # =========================================================
    ws5 = wb.create_sheet(title="05_PARAMETROS")
    ws5.views.sheetView[0].showGridLines = True

    params = {
        "TIPOS_MOVIMIENTO": ["Ingreso", "Gasto", "Transferencia"],
        "CATEGORIAS_GASTO": ["Alimentación", "Transporte/Moto", "Vivienda/Arriendo", "Servicios", "Salud/Gym", "Ocio/Restaurante", "Educación", "Compras", "Imprevistos"],
        "CATEGORIAS_INGRESO": ["Salario", "Honorarios", "Venta", "Rendimientos", "Préstamo Cobrado", "Otros"],
        "MEDIOS_PAGO": ["AyC Colanta", "Nequi", "Efectivo", "Bancolombia", "Daviplata", "Tarjeta Débito", "Tarjeta Crédito"]
    }

    col_p = 1
    for cat_title, items in params.items():
        ws5.cell(row=2, column=col_p, value=cat_title).font = font_header
        ws5.cell(row=2, column=col_p).fill = fill_header
        ws5.cell(row=2, column=col_p).alignment = align_center
        ws5.cell(row=2, column=col_p).border = thin_border

        for idx, item in enumerate(items, 3):
            c = ws5.cell(row=idx, column=col_p, value=item)
            c.font = font_body
            c.alignment = align_left
            c.border = thin_border

        col_p += 2

    # Auto-ajustar anchos de columnas
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if not val.startswith("=") and len(val) > max_len and len(val) < 60:
                    max_len = len(val)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(filename)
    print(f"Master multi-user Excel successfully generated at: {filename}")

if __name__ == "__main__":
    build_multiuser_excel()

import sys
if hasattr(sys.stdout, "reconfigure"): sys.stdout.reconfigure(encoding="utf-8")
import openpyxl

wb = openpyxl.load_workbook('Panel_Finanzas.xlsx', data_only=True)
for sheet in ['PANEL', 'MOVIMIENTOS', 'CONTROL FIJOS', 'AHORROS']:
    if sheet not in wb.sheetnames:
        continue
    ws = wb[sheet]
    print(f"\n==================== SHEET: {sheet} ====================")
    for r in range(1, 35):
        vals = [ws.cell(r, c).value for c in range(1, 16)]
        if any(v is not None for v in vals):
            row_str = " | ".join(str(v) if v is not None else "" for v in vals[:14])
            print(f"R{r:02d}: {row_str}")

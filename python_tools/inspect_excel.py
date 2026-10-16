import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
import openpyxl
import json

wb = openpyxl.load_workbook('Panel_Finanzas.xlsx', data_only=True)
print("=== SHEETS IN Panel_Finanzas.xlsx ===")
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\nSheet: '{name}' | Rows: {ws.max_row} | Cols: {ws.max_column}")
    for r in range(1, min(15, ws.max_row + 1)):
        vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
        if any(v is not None for v in vals):
            # Print non-null or formatted
            formatted = [str(v) if v is not None else "" for v in vals]
            print(f"  R{r}: {formatted[:10]}")
